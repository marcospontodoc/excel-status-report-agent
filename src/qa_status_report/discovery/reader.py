"""Read-only Excel workbook loading and structural header discovery."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
import hashlib
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError


class WorkbookDiscoveryError(ValueError):
    """Base error for invalid or unreadable workbook inputs."""


class UnsupportedWorkbookFormatError(WorkbookDiscoveryError):
    """Raised when the path is not an .xlsx or .xlsm workbook."""


class CorruptWorkbookError(WorkbookDiscoveryError):
    """Raised when an Excel workbook cannot be parsed."""


@dataclass
class ColumnData:
    index: int
    letter: str
    header: str | None
    header_row: int | None
    observed_types: dict[str, int]
    non_empty_count: int
    missing_count: int
    distinct_value_hashes: set[str]
    formula_observations: list[tuple[int, bool, str | None]]


@dataclass
class SheetData:
    name: str
    header_row: int | None
    columns: list[ColumnData]
    data_row_count: int
    empty: bool


@dataclass
class WorkbookData:
    path: Path
    sheets: list[SheetData]
    limitations: list[str]


def _normalize_header(value: Any) -> str:
    return " ".join(str(value).strip().casefold().split()) if value is not None else ""


def _header_score(row: tuple[Any, ...]) -> float | None:
    values = [value for value in row if value is not None and str(value).strip()]
    if len(values) < 2:
        return None
    text_count = sum(isinstance(value, str) for value in values)
    distinct_ratio = len({_normalize_header(value) for value in values}) / len(values)
    text_ratio = text_count / len(values)
    if text_ratio < 0.6 or distinct_ratio < 0.6:
        return None
    return text_ratio * 0.6 + distinct_ratio * 0.4 + min(len(values), 8) * 0.01


def _detect_header(rows: list[tuple[Any, ...]]) -> int | None:
    candidates = [(_header_score(row), index + 1) for index, row in enumerate(rows[:20])]
    valid = [(score, row_number) for score, row_number in candidates if score is not None]
    if not valid:
        return None
    # Stable tie-breaking favors the earliest plausible header row.
    return max(valid, key=lambda item: (item[0], -item[1]))[1]


def _type_name(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, (date, time)):
        return "date" if isinstance(value, date) else "time"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "decimal"
    if isinstance(value, str):
        return "string"
    return type(value).__name__


def _value_digest(value: Any) -> str:
    if isinstance(value, str):
        normalized = value.strip().casefold()
    elif isinstance(value, (date, datetime, time)):
        normalized = value.isoformat()
    else:
        normalized = str(value)
    material = f"{type(value).__name__}:{normalized}".encode("utf-8", errors="replace")
    return hashlib.sha256(material).hexdigest()


def read_workbook(path: str | Path) -> WorkbookData:
    """Read .xlsx/.xlsm in read-only mode; never writes to the input file."""
    workbook_path = Path(path).expanduser()
    if not workbook_path.exists() or not workbook_path.is_file():
        raise FileNotFoundError(f"Arquivo Excel não encontrado: {workbook_path}")
    suffix = workbook_path.suffix.casefold()
    if suffix not in {".xlsx", ".xlsm"}:
        raise UnsupportedWorkbookFormatError(
            f"Formato não suportado: {suffix or '(sem extensão)'}. Formatos aceitos: .xlsx e .xlsm."
        )

    formulas_book = None
    cached_book = None
    try:
        formulas_book = load_workbook(workbook_path, read_only=True, data_only=False)
        cached_book = load_workbook(workbook_path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError, KeyError) as exc:
        if formulas_book is not None:
            formulas_book.close()
        if cached_book is not None:
            cached_book.close()
        raise CorruptWorkbookError(f"Não foi possível ler o arquivo Excel: {workbook_path.name}") from exc

    sheets: list[SheetData] = []
    limitations: list[str] = []
    try:
        for formula_sheet in formulas_book.worksheets:
            cached_sheet = cached_book[formula_sheet.title]
            initial_rows = list(formula_sheet.iter_rows(min_row=1, max_row=min(formula_sheet.max_row or 1, 20), values_only=True))
            header_row = _detect_header(initial_rows)
            max_column = formula_sheet.max_column or 0
            if max_column == 0 or (formula_sheet.max_row or 0) == 0:
                sheets.append(SheetData(formula_sheet.title, header_row, [], 0, True))
                continue

            headers: dict[int, str | None] = {}
            if header_row is not None and header_row <= len(initial_rows):
                for column_index, value in enumerate(initial_rows[header_row - 1], start=1):
                    if column_index <= max_column:
                        headers[column_index] = str(value).strip() if value is not None and str(value).strip() else None

            columns = [
                ColumnData(index, get_column_letter(index), headers.get(index), header_row, {}, 0, 0, set(), [])
                for index in range(1, max_column + 1)
            ]
            data_start = (header_row + 1) if header_row is not None else 1
            data_row_count = max(0, (formula_sheet.max_row or 0) - data_start + 1)
            if data_row_count:
                formula_rows = formula_sheet.iter_rows(min_row=data_start, max_col=max_column, values_only=False)
                cached_rows = cached_sheet.iter_rows(min_row=data_start, max_col=max_column, values_only=False)
                for row_number, (formula_row, cached_row) in enumerate(zip(formula_rows, cached_rows), start=data_start):
                    for index, (formula_cell, cached_cell) in enumerate(zip(formula_row, cached_row), start=1):
                        is_formula = formula_cell.data_type == "f"
                        value = cached_cell.value if is_formula else formula_cell.value
                        column = columns[index - 1]
                        if is_formula:
                            cached_available = value is not None
                            cached_type = _type_name(value) if cached_available else None
                            column.formula_observations.append((row_number, cached_available, cached_type))
                            if not cached_available:
                                limitations.append(
                                    f"Fórmula sem valor em cache detectável em {formula_sheet.title}!{formula_cell.coordinate}."
                                )
                        if is_formula and value is None:
                            continue
                        if value is None or (isinstance(value, str) and not value.strip()):
                            column.missing_count += 1
                        else:
                            observed_type = _type_name(value)
                            if observed_type:
                                column.observed_types[observed_type] = column.observed_types.get(observed_type, 0) + 1
                            column.non_empty_count += 1
                            column.distinct_value_hashes.add(_value_digest(value))
            has_data = any(column.non_empty_count or column.formula_observations for column in columns)
            sheets.append(SheetData(formula_sheet.title, header_row, columns, data_row_count, not has_data and header_row is None))
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError, KeyError) as exc:
        raise CorruptWorkbookError(f"Não foi possível processar o arquivo Excel: {workbook_path.name}") from exc
    finally:
        formulas_book.close()
        cached_book.close()

    return WorkbookData(workbook_path, sheets, sorted(set(limitations)))


__all__ = [
    "CorruptWorkbookError",
    "UnsupportedWorkbookFormatError",
    "WorkbookDiscoveryError",
    "WorkbookData",
    "read_workbook",
]
