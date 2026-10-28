"""Read and validate the project's fixed Excel workbook without modifying it."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from zipfile import BadZipFile

from .fixed_models import (
    Bug,
    CellOrigin,
    Functionality,
    ProjectInfo,
    QAWorkbook,
    Scenario,
    WorkbookNote,
)

SHEET_HEADERS: dict[str, tuple[str, ...]] = {
    "Projeto": ("Nome do Projeto", "Início do Período", "Fim do Período"),
    "Funcionalidades": ("ID da Funcionalidade", "Nome da Funcionalidade"),
    "Cenários": (
        "ID da Funcionalidade", "ID do Cenário", "Descrição do Cenário",
        "Pré-requisitos", "Passo a passo", "Status", "Responsável", "Observações",
    ),
    "Bugs": (
        "ID do Bug", "ID da Funcionalidade", "ID do Cenário", "Descrição do Bug",
        "Status", "Severidade", "Responsável", "Data de Abertura",
    ),
}
SCENARIO_STATUSES = ("Falha", "Ignorado", "Bloqueado", "Em Andamento", "Aprovado")


class WorkbookInputError(ValueError):
    """Invalid workbook or data that cannot safely support a report."""


def _origin(filename: str, sheet: str, column: str, row: int, column_index: int) -> CellOrigin:
    return CellOrigin(filename, sheet, column, row, f"{get_column_letter(column_index)}{row}")


def _empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _identifier(value: Any, label: str, origin: CellOrigin, *, is_formula: bool = False) -> str:
    if is_formula and value is None:
        raise WorkbookInputError(f"Fórmula sem valor em cache em {origin.sheet}!{origin.coordinate}; {label} não pode ser inferido e a fórmula não foi recalculada.")
    if isinstance(value, bool) or not isinstance(value, (str, int)):
        raise WorkbookInputError(f"{label} inválido em {origin.sheet}!{origin.coordinate}; informe texto ou número inteiro.")
    text = str(value).strip()
    if not text:
        raise WorkbookInputError(f"{label} obrigatório ausente em {origin.sheet}!{origin.coordinate}.")
    return text


def _text(value: Any, label: str, origin: CellOrigin, *, required: bool, is_formula: bool = False) -> str | None:
    if required and is_formula and value is None:
        raise WorkbookInputError(f"Fórmula sem valor em cache em {origin.sheet}!{origin.coordinate}; {label} não pode ser inferido e a fórmula não foi recalculada.")
    if _empty(value):
        if required:
            raise WorkbookInputError(f"{label} obrigatório ausente em {origin.sheet}!{origin.coordinate}.")
        return None
    if not isinstance(value, str):
        raise WorkbookInputError(f"{label} deve ser texto em {origin.sheet}!{origin.coordinate}.")
    return value.strip()


def _date(value: Any, label: str, origin: CellOrigin, *, is_formula: bool = False) -> date:
    if is_formula and value is None:
        raise WorkbookInputError(f"Fórmula sem valor em cache em {origin.sheet}!{origin.coordinate}; {label} não pode ser inferido e a fórmula não foi recalculada.")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                pass
    raise WorkbookInputError(f"{label} deve ser uma data válida em {origin.sheet}!{origin.coordinate}.")


def read_fixed_workbook(path: str | Path) -> QAWorkbook:
    """Load the fixed workbook and reject structural or referential errors.

    Workbooks are opened read-only. Formula cells are never recalculated; their
    cached values are used when present and missing caches are reported or block
    a required field.
    """
    workbook_path = Path(path).expanduser()
    if not workbook_path.exists() or not workbook_path.is_file():
        raise FileNotFoundError(f"Arquivo Excel não encontrado: {workbook_path}")
    if workbook_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise WorkbookInputError("Formato não suportado. Use um arquivo .xlsx ou .xlsm.")

    formula_book = cached_book = None
    notes: list[WorkbookNote] = []
    errors: list[str] = []
    try:
        formula_book = load_workbook(workbook_path, read_only=True, data_only=False, keep_vba=workbook_path.suffix.casefold() == ".xlsm")
        cached_book = load_workbook(workbook_path, read_only=True, data_only=True, keep_vba=workbook_path.suffix.casefold() == ".xlsm")
        required_sheets = tuple(SHEET_HEADERS)
        missing_sheets = [name for name in required_sheets if name not in formula_book.sheetnames]
        extra_sheets = [name for name in formula_book.sheetnames if name not in SHEET_HEADERS]
        if missing_sheets:
            errors.append("Abas obrigatórias ausentes: " + ", ".join(missing_sheets) + ".")
        if extra_sheets:
            errors.append("Abas não previstas no modelo fixo: " + ", ".join(extra_sheets) + ".")
        if errors:
            raise WorkbookInputError("\n".join(errors))

        headers_by_sheet: dict[str, dict[str, int]] = {}
        for sheet_name, expected in SHEET_HEADERS.items():
            sheet = formula_book[sheet_name]
            header_cells = next(sheet.iter_rows(min_row=1, max_row=1, values_only=False), ())
            header_map: dict[str, int] = {}
            duplicates: set[str] = set()
            for column_index, cell in enumerate(header_cells, start=1):
                value = cell.value
                if isinstance(value, str) and value.strip():
                    header = value.strip()
                    if header in header_map:
                        duplicates.add(header)
                    else:
                        header_map[header] = column_index
            missing = [name for name in expected if name not in header_map]
            if missing:
                errors.append(f"{sheet_name}: cabeçalhos obrigatórios ausentes: {', '.join(missing)}.")
            if duplicates:
                errors.append(f"{sheet_name}: cabeçalhos duplicados: {', '.join(sorted(duplicates))}.")
            extra = [name for name in header_map if name not in expected]
            if extra:
                notes.append(WorkbookNote("extra_columns", f"Colunas adicionais ignoradas em {sheet_name}: {', '.join(extra)}."))
            headers_by_sheet[sheet_name] = header_map
        if errors:
            raise WorkbookInputError("\n".join(errors))

        def get_value(sheet_name: str, row_number: int, column_name: str) -> tuple[Any, CellOrigin, bool]:
            column_index = headers_by_sheet[sheet_name][column_name]
            origin = _origin(str(workbook_path.resolve()), sheet_name, column_name, row_number, column_index)
            source_cell = formula_book[sheet_name].cell(row=row_number, column=column_index)
            value_cell = cached_book[sheet_name].cell(row=row_number, column=column_index)
            is_formula = source_cell.data_type == "f"
            value = value_cell.value if is_formula else source_cell.value
            if is_formula:
                if value is None:
                    notes.append(WorkbookNote("formula_without_cache", "Fórmula sem valor calculado em cache; nenhuma fórmula foi recalculada.", origin))
                else:
                    notes.append(WorkbookNote("formula_cache_used", "Valor em cache de fórmula utilizado sem recalcular a fórmula; o cache pode estar desatualizado.", origin))
            return value, origin, is_formula

        def _is_formula(sheet_name: str, row_number: int, column_name: str) -> bool:
            column_index = headers_by_sheet[sheet_name][column_name]
            return formula_book[sheet_name].cell(row=row_number, column=column_index).data_type == "f"

        def populated_rows(sheet_name: str) -> list[int]:
            sheet = formula_book[sheet_name]
            columns = tuple(headers_by_sheet[sheet_name].values())
            rows: list[int] = []
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), start=2):
                if any(not _empty(row[index - 1].value) for index in columns if index <= len(row)):
                    rows.append(row_number)
            return rows

        project_rows = populated_rows("Projeto")
        if len(project_rows) != 1:
            raise WorkbookInputError(f"A aba Projeto deve conter exatamente um registro preenchido; encontrados {len(project_rows)}.")
        project_row = project_rows[0]
        name_val, name_origin, _ = get_value("Projeto", project_row, "Nome do Projeto")
        start_val, start_origin, _ = get_value("Projeto", project_row, "Início do Período")
        end_val, end_origin, _ = get_value("Projeto", project_row, "Fim do Período")
        project_name = _text(name_val, "Nome do Projeto", name_origin, required=True, is_formula=_is_formula("Projeto", project_row, "Nome do Projeto"))
        assert project_name is not None
        project = ProjectInfo(
            project_name,
            _date(start_val, "Início do Período", start_origin, is_formula=_is_formula("Projeto", project_row, "Início do Período")),
            _date(end_val, "Fim do Período", end_origin, is_formula=_is_formula("Projeto", project_row, "Fim do Período")),
            (name_origin, start_origin, end_origin),
        )
        if project.period_end < project.period_start:
            raise WorkbookInputError("Fim do período anterior ao início em Projeto; confira as datas.")

        functionalities: list[Functionality] = []
        functionality_ids: set[str] = set()
        for row_number in populated_rows("Funcionalidades"):
            id_value, id_origin, _ = get_value("Funcionalidades", row_number, "ID da Funcionalidade")
            name_value, name_origin, _ = get_value("Funcionalidades", row_number, "Nome da Funcionalidade")
            functionality_id = _identifier(id_value, "ID da Funcionalidade", id_origin, is_formula=_is_formula("Funcionalidades", row_number, "ID da Funcionalidade"))
            functionality_name = _text(name_value, "Nome da Funcionalidade", name_origin, required=True, is_formula=_is_formula("Funcionalidades", row_number, "Nome da Funcionalidade"))
            assert functionality_name is not None
            if functionality_id in functionality_ids:
                raise WorkbookInputError(f"ID de funcionalidade duplicado {functionality_id!r} em {id_origin.sheet}!{id_origin.coordinate}.")
            functionality_ids.add(functionality_id)
            functionalities.append(Functionality(functionality_id, functionality_name, (id_origin, name_origin)))

        scenarios: list[Scenario] = []
        scenario_by_id: dict[str, Scenario] = {}
        for row_number in populated_rows("Cenários"):
            cols = {name: get_value("Cenários", row_number, name) for name in SHEET_HEADERS["Cenários"]}
            f_id = _identifier(cols["ID da Funcionalidade"][0], "ID da Funcionalidade", cols["ID da Funcionalidade"][1], is_formula=cols["ID da Funcionalidade"][2])
            scenario_id = _identifier(cols["ID do Cenário"][0], "ID do Cenário", cols["ID do Cenário"][1], is_formula=cols["ID do Cenário"][2])
            description = _text(cols["Descrição do Cenário"][0], "Descrição do Cenário", cols["Descrição do Cenário"][1], required=True, is_formula=cols["Descrição do Cenário"][2])
            status_raw = cols["Status"][0]
            if cols["Status"][2] and status_raw is None:
                raise WorkbookInputError(f"Fórmula sem valor em cache em Cenários!{cols['Status'][1].coordinate}; o status do cenário não pode ser inferido.")
            status_blank = _empty(status_raw)
            if status_blank:
                status = "Em Andamento"
            elif isinstance(status_raw, str) and status_raw.strip() in SCENARIO_STATUSES:
                status = status_raw.strip()
            else:
                raise WorkbookInputError(f"Status inválido em Cenários!{cols['Status'][1].coordinate}; valores permitidos: {', '.join(SCENARIO_STATUSES)}.")
            responsible = _text(cols["Responsável"][0], "Responsável", cols["Responsável"][1], required=False, is_formula=cols["Responsável"][2])
            observation = _text(cols["Observações"][0], "Observações", cols["Observações"][1], required=False, is_formula=cols["Observações"][2])
            if f_id not in functionality_ids:
                raise WorkbookInputError(f"Funcionalidade {f_id!r} não encontrada para o cenário {scenario_id!r} ({cols['ID da Funcionalidade'][1].sheet}!{cols['ID da Funcionalidade'][1].coordinate}).")
            if scenario_id in scenario_by_id:
                raise WorkbookInputError(f"ID de cenário duplicado {scenario_id!r} em {cols['ID do Cenário'][1].sheet}!{cols['ID do Cenário'][1].coordinate}.")
            assert description is not None
            origins = tuple(item[1] for item in cols.values())
            scenario = Scenario(scenario_id, f_id, description, status, responsible, observation, origins, status_blank)
            scenario_by_id[scenario_id] = scenario
            scenarios.append(scenario)

        bugs: list[Bug] = []
        bug_ids: set[str] = set()
        for row_number in populated_rows("Bugs"):
            cols = {name: get_value("Bugs", row_number, name) for name in SHEET_HEADERS["Bugs"]}
            bug_id = _identifier(cols["ID do Bug"][0], "ID do Bug", cols["ID do Bug"][1], is_formula=cols["ID do Bug"][2])
            f_id = _identifier(cols["ID da Funcionalidade"][0], "ID da Funcionalidade", cols["ID da Funcionalidade"][1], is_formula=cols["ID da Funcionalidade"][2])
            scenario_id = _identifier(cols["ID do Cenário"][0], "ID do Cenário", cols["ID do Cenário"][1], is_formula=cols["ID do Cenário"][2])
            description = _text(cols["Descrição do Bug"][0], "Descrição do Bug", cols["Descrição do Bug"][1], required=True, is_formula=cols["Descrição do Bug"][2])
            status = _text(cols["Status"][0], "Status do bug", cols["Status"][1], required=True, is_formula=cols["Status"][2])
            severity = _text(cols["Severidade"][0], "Severidade", cols["Severidade"][1], required=True, is_formula=cols["Severidade"][2])
            responsible = _text(cols["Responsável"][0], "Responsável do bug", cols["Responsável"][1], required=True, is_formula=cols["Responsável"][2])
            opened_at = _date(cols["Data de Abertura"][0], "Data de Abertura", cols["Data de Abertura"][1], is_formula=cols["Data de Abertura"][2])
            if bug_id in bug_ids:
                raise WorkbookInputError(f"ID de bug duplicado {bug_id!r} em {cols['ID do Bug'][1].sheet}!{cols['ID do Bug'][1].coordinate}.")
            bug_ids.add(bug_id)
            if f_id not in functionality_ids:
                raise WorkbookInputError(f"Funcionalidade {f_id!r} não encontrada para o bug {bug_id!r}.")
            scenario = scenario_by_id.get(scenario_id)
            if scenario is None:
                raise WorkbookInputError(f"Cenário {scenario_id!r} não encontrado para o bug {bug_id!r}.")
            if scenario.functionality_id != f_id:
                raise WorkbookInputError(f"Vínculo inconsistente no bug {bug_id!r}: o cenário {scenario_id!r} pertence a {scenario.functionality_id!r}, não a {f_id!r}.")
            assert description is not None and status is not None and severity is not None and responsible is not None
            bugs.append(Bug(bug_id, f_id, scenario_id, description, status, severity, responsible, opened_at, tuple(item[1] for item in cols.values())))

        return QAWorkbook(str(workbook_path.resolve()), project, tuple(functionalities), tuple(scenarios), tuple(bugs), tuple(notes))
    except (BadZipFile, InvalidFileException, OSError, KeyError, IndexError) as exc:
        raise WorkbookInputError(f"Arquivo Excel inválido ou corrompido: {workbook_path.name}.") from exc
    finally:
        if formula_book is not None:
            formula_book.close()
        if cached_book is not None:
            cached_book.close()


__all__ = ["SCENARIO_STATUSES", "SHEET_HEADERS", "WorkbookInputError", "read_fixed_workbook"]
