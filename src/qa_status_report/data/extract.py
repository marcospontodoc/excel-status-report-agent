"""Extract canonical records from a workbook using only explicit decisions."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError

from qa_status_report.discovery.models import DiscoveryResult, SourceRef
from qa_status_report.discovery.reader import (
    CorruptWorkbookError,
    UnsupportedWorkbookFormatError,
)

from .models import (
    CanonicalDataset,
    CanonicalRecord,
    EntityName,
    ExtractionResult,
    MappingDecision,
    MappingDecisionSet,
    Provenance,
    RecordValue,
    ValidationIssue,
    ValidationState,
)

_FIELD_TYPES: dict[str, str] = {
    "project.id": "identifier", "project.name": "text",
    "project.reference_period_start": "date", "project.reference_period_end": "date",
    "sprint.id": "identifier", "sprint.project_id": "reference", "sprint.name": "text",
    "sprint.start_date": "date", "sprint.end_date": "date",
    "release.id": "identifier", "release.project_id": "reference", "release.name": "text",
    "release.start_date": "date", "release.target_date": "date", "release.released_at": "date",
    "requirement.id": "identifier", "requirement.project_id": "reference",
    "requirement.title": "text", "requirement.status_raw": "raw_value",
    "test_scenario.id": "identifier", "test_scenario.project_id": "reference",
    "test_scenario.title": "text", "test_scenario.requirement_ids": "reference_list",
    "test_execution.id": "identifier", "test_execution.scenario_id": "reference",
    "test_execution.executed_at": "datetime", "test_execution.sprint_id": "reference",
    "test_execution.release_id": "reference", "test_execution.status_raw": "raw_value",
    "test_execution.outcome": "text", "test_execution.attempt_number": "integer",
    "test_execution.retest_of_id": "reference", "bug.id": "identifier",
    "bug.project_id": "reference", "bug.title": "text", "bug.status_raw": "raw_value",
    "bug.severity_raw": "raw_value", "bug.priority_raw": "raw_value",
    "bug.release_ids": "reference_list", "bug.scenario_ids": "reference_list",
    "bug.execution_ids": "reference_list",
}
_ENTITY_NAMES = {field.split(".", 1)[0] for field in _FIELD_TYPES}


def _provenance(source: SourceRef) -> Provenance:
    if source.row is None or source.coordinate is None or source.column is None:
        raise ValueError("Decisão confirmada precisa referenciar aba, coluna e linha/célula de origem.")
    return Provenance(source.workbook, source.sheet, source.column, source.row, source.coordinate)


def _field_entity(canonical_field: str) -> EntityName:
    entity = canonical_field.split(".", 1)[0]
    if entity not in _ENTITY_NAMES or canonical_field not in _FIELD_TYPES:
        raise ValueError(f"Campo canônico desconhecido: {canonical_field}")
    return entity  # type: ignore[return-value]


def _is_valid_type(value: Any, canonical_type: str) -> bool:
    if canonical_type in {"identifier", "reference"}:
        return isinstance(value, (str, int)) and not isinstance(value, bool)
    if canonical_type == "text":
        return isinstance(value, str)
    if canonical_type == "date":
        return isinstance(value, (date, datetime))
    if canonical_type == "datetime":
        return isinstance(value, datetime)
    if canonical_type == "integer":
        return isinstance(value, int) and not isinstance(value, bool)
    if canonical_type == "raw_value":
        return isinstance(value, (str, int, float, bool, date, datetime, time))
    # No list delimiter/escaping rule has been approved; retain but do not normalize.
    return False


def _validated_mapping_sources(
    path: Path,
    discovery: DiscoveryResult,
    decision_set: MappingDecisionSet,
) -> tuple[list[tuple[MappingDecision, Provenance]], list[ValidationIssue]]:
    issues: list[ValidationIssue] = []
    confirmed: list[tuple[MappingDecision, Provenance]] = []
    seen_source_fields: set[tuple[str, str]] = set()
    profile_sheets = {sheet.name: sheet for sheet in discovery.sheets}
    for decision in decision_set.decisions:
        if decision.status not in {"confirmed", "ignored", "pending"}:
            issues.append(ValidationIssue("invalid_decision_status", "Estado de decisão de mapeamento desconhecido.", decision.canonical_field, decision.source, True))
            continue
        try:
            entity = _field_entity(decision.canonical_field)
        except ValueError as exc:
            issues.append(ValidationIssue("unknown_canonical_field", str(exc), decision.canonical_field, blocking=True))
            continue
        if decision.status != "confirmed":
            if decision.status == "pending" and (decision.required or decision.blocking):
                continue
            if decision.status == "ignored" and decision.required:
                continue
            if decision.status == "confirmed":
                continue
            if decision.source is not None:
                issues.append(ValidationIssue("nonconfirmed_mapping_has_source", "Mapeamento não confirmado não pode fornecer origem para extração.", decision.canonical_field, decision.source, decision.blocking))
            continue
        if decision.source is None:
            issues.append(ValidationIssue("confirmed_mapping_without_source", "Mapeamento confirmado precisa de origem explícita.", decision.canonical_field, blocking=True))
            continue
        source = decision.source
        if Path(source.workbook).name != path.name or decision_set.workbook != path.name or discovery.workbook != path.name:
            issues.append(ValidationIssue("workbook_mismatch", "Decisão, descoberta e arquivo precisam identificar a mesma pasta de trabalho.", decision.canonical_field, source, True))
            continue
        sheet_profile = profile_sheets.get(source.sheet)
        if sheet_profile is None:
            issues.append(ValidationIssue("source_sheet_not_discovered", "A aba da decisão não existe na descoberta fornecida.", decision.canonical_field, source, True))
            continue
        if source.row != sheet_profile.header_row:
            issues.append(ValidationIssue("source_header_row_mismatch", "A linha de cabeçalho da decisão difere da descoberta.", decision.canonical_field, source, True))
            continue
        try:
            column_letter, _ = coordinate_from_string(source.coordinate)
            column_index = column_index_from_string(column_letter)
        except (TypeError, ValueError):
            issues.append(ValidationIssue("invalid_source_coordinate", "A coordenada de origem não é válida.", decision.canonical_field, source, True))
            continue
        profiled_column = next(
            (
                column
                for column in sheet_profile.columns
                if column.source.coordinate
                and coordinate_from_string(column.source.coordinate)[0] == column_letter
            ),
            None,
        )
        if profiled_column is None or profiled_column.source.column != source.column:
            issues.append(ValidationIssue("source_column_mismatch", "A aba/coluna da decisão não corresponde ao perfil descoberto.", decision.canonical_field, source, True))
            continue
        source_key = (source.sheet, decision.canonical_field)
        if source_key in seen_source_fields:
            issues.append(ValidationIssue("duplicate_mapping", "Campo canônico mapeado mais de uma vez na mesma aba.", decision.canonical_field, source, True))
            continue
        seen_source_fields.add(source_key)
        confirmed.append((decision, _provenance(source)))
    return confirmed, issues


def extract_records(
    path: str | Path,
    discovery: DiscoveryResult,
    decisions: MappingDecisionSet,
) -> ExtractionResult:
    """Read records and create a calculable dataset only after the global gate passes."""
    workbook_path = Path(path).expanduser()
    if not workbook_path.exists() or not workbook_path.is_file():
        raise FileNotFoundError(f"Arquivo Excel não encontrado: {workbook_path}")
    if workbook_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise UnsupportedWorkbookFormatError("Formato não suportado. Formatos aceitos: .xlsx e .xlsm.")
    if discovery.workbook != workbook_path.name or decisions.workbook != workbook_path.name:
        raise ValueError("DiscoveryResult e decisões precisam pertencer ao arquivo informado.")

    issues: list[ValidationIssue] = []
    confirmed, mapping_issues = _validated_mapping_sources(workbook_path, discovery, decisions)
    issues.extend(mapping_issues)
    if not decisions.review_complete:
        issues.append(ValidationIssue("mapping_review_incomplete", "A revisão consolidada dos mapeamentos ainda não foi concluída.", blocking=True))
    for ambiguity in decisions.blocking_ambiguities:
        issues.append(ValidationIssue("blocking_ambiguity", ambiguity, blocking=True))

    confirmed_fields = {decision.canonical_field for decision, _ in confirmed}
    pending_required = set(decisions.required_fields)
    pending_required.update(
        decision.canonical_field
        for decision in decisions.decisions
        if decision.required and decision.status != "confirmed"
    )
    pending_required.difference_update(confirmed_fields)
    for field_name in sorted(pending_required):
        issues.append(ValidationIssue("required_mapping_pending", f"Decisão obrigatória pendente: {field_name}.", field_name, blocking=True))
    for decision in decisions.decisions:
        if decision.status == "pending" and decision.blocking:
            issues.append(ValidationIssue("blocking_mapping_pending", f"Mapeamento bloqueante pendente: {decision.canonical_field}.", decision.canonical_field, decision.source, True))
    if any(issue.blocking for issue in issues):
        confirmed = []

    formula_book = cached_book = None
    try:
        formula_book = load_workbook(workbook_path, read_only=True, data_only=False)
        cached_book = load_workbook(workbook_path, read_only=True, data_only=True)
        for decision, source in confirmed:
            if source.sheet not in formula_book.sheetnames:
                issues.append(
                    ValidationIssue(
                        "source_sheet_changed",
                        "A aba da origem confirmada não existe no arquivo atual; refaça a descoberta/revisão.",
                        decision.canonical_field,
                        source,
                        True,
                    )
                )
                continue
            letter, _ = coordinate_from_string(source.coordinate)
            source_sheet = formula_book[source.sheet]
            header_value = source_sheet.cell(row=source.row, column=column_index_from_string(letter)).value
            actual_header = str(header_value).strip() if header_value is not None and str(header_value).strip() else letter
            if actual_header != source.column:
                issues.append(
                    ValidationIssue(
                        "source_header_changed",
                        "O cabeçalho no arquivo atual difere da origem confirmada; refaça a descoberta/revisão.",
                        decision.canonical_field,
                        source,
                        True,
                    )
                )
        if any(issue.blocking for issue in issues):
            confirmed = []
        grouped: dict[tuple[str, str], list[tuple[MappingDecision, Provenance, int]]] = defaultdict(list)
        for decision, source in confirmed:
            letter, _ = coordinate_from_string(source.coordinate)
            grouped[(_field_entity(decision.canonical_field), source.sheet)].append(
                (decision, source, column_index_from_string(letter))
            )

        records: list[CanonicalRecord] = []
        limitations: list[str] = []
        for (entity, sheet_name), field_mappings in grouped.items():
            if sheet_name not in formula_book.sheetnames:
                issues.append(ValidationIssue("source_sheet_missing", f"Aba não encontrada durante extração: {sheet_name}.", blocking=True))
                continue
            ws_formula = formula_book[sheet_name]
            ws_cached = cached_book[sheet_name]
            header_row = field_mappings[0][1].row
            if any(source.row != header_row for _, source, _ in field_mappings):
                issues.append(ValidationIssue("inconsistent_header_rows", f"Mapeamentos para {entity} na aba {sheet_name} usam linhas de cabeçalho diferentes.", blocking=True))
                continue
            max_column = max(ws_formula.max_column or 0, *(index for _, _, index in field_mappings))
            formula_rows = ws_formula.iter_rows(min_row=header_row + 1, max_col=max_column, values_only=False)
            cached_rows = ws_cached.iter_rows(min_row=header_row + 1, max_col=max_column, values_only=False)
            for row_number, (formula_row, cached_row) in enumerate(zip(formula_rows, cached_rows), start=header_row + 1):
                if not any(cell.value is not None for cell in formula_row):
                    continue
                fields: dict[str, RecordValue] = {}
                record_provenance: list[Provenance] = []
                for decision, header_source, column_index in field_mappings:
                    formula_cell = formula_row[column_index - 1]
                    cached_cell = cached_row[column_index - 1]
                    is_formula = formula_cell.data_type == "f"
                    value = cached_cell.value if is_formula else formula_cell.value
                    coordinate = f"{get_column_letter(column_index)}{row_number}"
                    cell_source = Provenance(
                        workbook_path.name,
                        sheet_name,
                        header_source.column,
                        row_number,
                        coordinate,
                    )
                    record_provenance.append(cell_source)
                    if is_formula and value is None:
                        field_value = RecordValue("formula_without_cache", provenance=cell_source, is_formula=True, cached_value_available=False, issue="Valor calculado em cache não disponível; fórmula não recalculada.")
                        limitations.append(f"Fórmula sem cache em {sheet_name}!{coordinate}.")
                    elif value is None or (isinstance(value, str) and not value.strip()):
                        field_value = RecordValue("missing", provenance=cell_source, is_formula=is_formula, cached_value_available=(value is not None if is_formula else None))
                    elif not _is_valid_type(value, _FIELD_TYPES[decision.canonical_field]):
                        message = "Formato de lista de referências não aprovado." if _FIELD_TYPES[decision.canonical_field] == "reference_list" else f"Tipo incompatível com {_FIELD_TYPES[decision.canonical_field]}."
                        field_value = RecordValue("invalid", raw_value=value, provenance=cell_source, is_formula=is_formula, cached_value_available=True if is_formula else None, issue=message)
                        issues.append(ValidationIssue("invalid_field_value", message, decision.canonical_field, cell_source, False))
                    else:
                        field_value = RecordValue("valid", value=value, raw_value=value, provenance=cell_source, is_formula=is_formula, cached_value_available=True if is_formula else None)
                    fields[decision.canonical_field.split(".", 1)[1]] = field_value
                records.append(CanonicalRecord(entity, f"{sheet_name}:{row_number}", sheet_name, fields, tuple(record_provenance)))
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError, KeyError) as exc:
        raise CorruptWorkbookError(f"Não foi possível processar o arquivo Excel: {workbook_path.name}") from exc
    finally:
        if formula_book is not None:
            formula_book.close()
        if cached_book is not None:
            cached_book.close()

    unresolved = tuple(decisions.blocking_ambiguities)
    pending_fields = tuple(sorted(pending_required))
    status = "blocked" if any(issue.blocking for issue in issues) else "ready"
    validation = ValidationState(status, tuple(issues), unresolved, pending_fields, decisions.review_complete)
    dataset = None
    if status == "ready":
        dataset = CanonicalDataset(workbook_path.name, tuple(records), decisions.decisions, validation)
    return ExtractionResult(workbook_path.name, tuple(records), validation, dataset, tuple(sorted(set(limitations))))


__all__ = ["extract_records"]
