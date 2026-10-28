"""Attach an editable Excel data snapshot to the generated native pie chart."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook

CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
PACKAGE_REL = f"{REL_NS}/package"
EMBEDDED_PATH = "ppt/embeddings/qa-status-chart-data.xlsx"
CHART_RELS = "ppt/slides/charts/_rels/chart1.xml.rels"

ET.register_namespace("c", CHART_NS)
ET.register_namespace("r", REL_NS)
ET.register_namespace("rel", PKG_REL_NS)
ET.register_namespace("", CONTENT_NS)


def _tag(namespace: str, name: str) -> str:
    return f"{{{namespace}}}{name}"


def _embedded_xlsx(distribution: tuple[tuple[str, int], ...]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QA Status"
    sheet.append(("Status", "Quantidade"))
    for status, count in distribution:
        sheet.append((status, count))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cache_ref(literal: ET.Element, *, ref_name: str, formula: str) -> ET.Element:
    result = ET.Element(_tag(CHART_NS, ref_name))
    f = ET.SubElement(result, _tag(CHART_NS, "f"))
    f.text = formula
    cache_name = "strCache" if ref_name == "strRef" else "numCache"
    cache = ET.SubElement(result, _tag(CHART_NS, cache_name))
    for child in literal:
        cache.append(deepcopy(child))
    return result


def _add_relationship(data: bytes) -> tuple[bytes, str]:
    if data:
        root = ET.fromstring(data)
    else:
        root = ET.Element(_tag(PKG_REL_NS, "Relationships"))
    existing_ids = {item.get("Id", "") for item in root}
    number = 1
    while f"rId{number}" in existing_ids:
        number += 1
    rel_id = f"rId{number}"
    ET.SubElement(root, _tag(PKG_REL_NS, "Relationship"), {
        "Id": rel_id,
        "Type": PACKAGE_REL,
        "Target": "../../embeddings/qa-status-chart-data.xlsx",
    })
    return ET.tostring(root, encoding="utf-8", xml_declaration=True), rel_id


def _add_content_type(data: bytes) -> bytes:
    root = ET.fromstring(data)
    part_name = f"/{EMBEDDED_PATH}"
    if not any(item.get("PartName") == part_name for item in root):
        ET.SubElement(root, _tag(CONTENT_NS, "Override"), {
            "PartName": part_name,
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        })
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def attach_editable_pie_data(path: str | Path, distribution: tuple[tuple[str, int], ...]) -> None:
    """Convert chart literals into cached worksheet references and embed the XLSX."""
    target = Path(path)
    if not distribution:
        raise ValueError("A distribuição de status precisa conter categorias para gerar o gráfico.")
    with ZipFile(target, "r") as source:
        chart_parts = [name for name in source.namelist() if name.startswith("ppt/") and "/charts/chart" in name and name.endswith(".xml")]
        if len(chart_parts) != 1:
            raise ValueError(f"Esperava um único gráfico nativo de pizza; encontrados {len(chart_parts)}.")
        chart_part = chart_parts[0]
        chart_rels = chart_part.rsplit("/", 1)[0] + "/_rels/" + chart_part.rsplit("/", 1)[1] + ".rels"
        archive_names = set(source.namelist())
        chart_root = ET.fromstring(source.read(chart_part))
        series = chart_root.findall(f".//{_tag(CHART_NS, 'ser')}")
        if len(series) != 1:
            raise ValueError(f"O gráfico deve ter uma série; encontradas {len(series)}.")
        series_node = series[0]
        replacements: list[tuple[ET.Element, ET.Element]] = []
        category_parent = series_node.find(f"{_tag(CHART_NS, 'cat')}")
        value_parent = series_node.find(f"{_tag(CHART_NS, 'val')}")
        category_literal = category_parent.find(f"{_tag(CHART_NS, 'strLit')}") if category_parent is not None else None
        value_literal = value_parent.find(f"{_tag(CHART_NS, 'numLit')}") if value_parent is not None else None
        if category_parent is None or value_parent is None or category_literal is None or value_literal is None:
            raise ValueError("O gráfico não contém dados literais de categorias e valores esperados.")
        categories_formula = f"'QA Status'!$A$2:$A${len(distribution) + 1}"
        values_formula = f"'QA Status'!$B$2:$B${len(distribution) + 1}"
        replacements.append((category_literal, _cache_ref(category_literal, ref_name="strRef", formula=categories_formula)))
        replacements.append((value_literal, _cache_ref(value_literal, ref_name="numRef", formula=values_formula)))
        for old, new in replacements:
            parent = old.getparent() if hasattr(old, "getparent") else None
            if parent is None:
                if old in category_parent:
                    category_parent.remove(old)
                    category_parent.append(new)
                elif old in value_parent:
                    value_parent.remove(old)
                    value_parent.append(new)
            else:  # pragma: no cover - ElementTree elements do not expose getparent.
                parent.replace(old, new)
        rel_data = source.read(chart_rels) if chart_rels in archive_names else b""
        new_rel_data, rel_id = _add_relationship(rel_data)
        external = ET.Element(_tag(CHART_NS, "externalData"), {_tag(REL_NS, "id"): rel_id})
        ET.SubElement(external, _tag(CHART_NS, "autoUpdate"), {"val": "0"})
        chart_root.append(external)
        chart_xml = ET.tostring(chart_root, encoding="utf-8", xml_declaration=True)
        content_types = _add_content_type(source.read("[Content_Types].xml"))
        embedded = _embedded_xlsx(distribution)

        output_buffer = BytesIO()
        with ZipFile(output_buffer, "w", compression=ZIP_DEFLATED) as output:
            for item in source.infolist():
                if item.filename == chart_part:
                    data = chart_xml
                elif item.filename == "[Content_Types].xml":
                    data = content_types
                elif item.filename == chart_rels:
                    data = new_rel_data
                else:
                    data = source.read(item.filename)
                output.writestr(item, data)
            if chart_rels not in archive_names:
                output.writestr(chart_rels, new_rel_data)
            output.writestr(EMBEDDED_PATH, embedded)
        final_bytes = output_buffer.getvalue()
        with ZipFile(BytesIO(final_bytes), "r") as check:
            if check.testzip() is not None:
                raise ValueError("O pacote PowerPoint gerado ficou corrompido ao anexar a planilha do gráfico.")
            workbook_check = load_workbook(BytesIO(check.read(EMBEDDED_PATH)), read_only=True, data_only=True)
            try:
                embedded_sheet = workbook_check["QA Status"]
                rows = list(embedded_sheet.iter_rows(min_row=2, values_only=True))
                if rows != list(distribution):
                    raise ValueError("A planilha editável do gráfico não corresponde às contagens calculadas.")
            finally:
                workbook_check.close()
    target.write_bytes(final_bytes)


__all__ = ["attach_editable_pie_data"]

